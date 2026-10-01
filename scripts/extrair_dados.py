"""Lê as planilhas de Dados/*.xlsx e atualiza data.js com todas as edições.

O data.js é a cópia de segurança do site (e a fonte, se a planilha do Google não estiver configurada).

Uso:
    py scripts/extrair_dados.py                 # todas as planilhas de Dados/
    py scripts/extrair_dados.py caminho.xlsx    # só uma planilha

Formatos aceitos:
  - Modelo MC (aba "Banco de Dados"), ex.: Dados/MC-2026.xlsx
      linha 1: SEMANA 01 (2 colunas por episódio) | linha 2: FOTO | COMPETIDOR | 1ª PROVA 1 | 2ª PROVA 1 ...
  - Formato antigo (aba "Placar"), ex.: Dados/TABELA MASTERCHEF 2026.xlsx

Cada planilha vira uma edição com código MC-AAAA (do nome do arquivo). Se duas planilhas derem o mesmo
código, vale a modificada por último. Nomes de provas e tabela de pontos que a planilha não tiver são
mantidos do data.js atual. Edições que não vieram de planilha (criadas pelo app) também são mantidas.
"""
import json
import re
import sys
import unicodedata
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parent.parent
DATA_JS = ROOT / "data.js"
PLACEHOLDER = re.compile(r"^\s*[12]\s*ª\s*PROVA(\s*\d+)?\s*$", re.I)
WEEK1 = re.compile(r"^SE[MN][AE]NA 0?1$", re.I)


def code_for(name_or_title):
    s = str(name_or_title).strip()
    if re.match(r"^MC-", s, re.I):
        return s.upper()
    y = re.search(r"\d{4}", s)
    if y:
        return "MC-" + y.group(0)
    t = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    return "MC-" + (re.sub(r"[^A-Za-z0-9]+", "-", t).strip("-").upper() or "EDICAO")


def title_for(code):
    return re.sub(r"^MC-?", "MasterChef ", code, flags=re.I).strip()


def tc(s):
    s = str(s or "").strip().lower()
    return re.sub(r"(^|\s|-)(\S)", lambda m: m.group(1) + m.group(2).upper(), s)


def week_label(raw, n):
    s = re.sub(r"^SEMENA", "SEMANA", str(raw or "").strip(), flags=re.I)
    return tc(s) if s else f"Semana {n:02d}"


def cell_code(v):
    s = str("" if v is None else v).strip().upper()
    if s in ("", "0", "·"):
        return "0"
    if s in ("–", "—"):
        return "-"
    return s


def find(ws, pattern):
    rx = re.compile(pattern, re.I) if isinstance(pattern, str) else pattern
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and rx.search(c.value.strip()):
                return c.row, c.column
    return None


def read_model(ws, comp, week1):
    """Modelo MC: COMPETIDOR na linha logo abaixo de SEMANA 01."""
    head_row, name_col, first_col = comp[0], comp[1], week1[1]
    weeks, col = [], first_col
    while True:
        label = ws.cell(week1[0], col).value
        has_data = any(str(ws.cell(r, col).value or "").strip() or str(ws.cell(r, col + 1).value or "").strip()
                       for r in range(head_row + 1, ws.max_row + 1))
        if not label and not has_data:
            break
        prov = [ws.cell(head_row, col).value, ws.cell(head_row, col + 1).value]
        weeks.append({"label": week_label(label, len(weeks) + 1),
                      "provas": ["" if (p is None or PLACEHOLDER.match(str(p))) else tc(p) for p in prov]})
        col += 2
    competitors = []
    for r in range(head_row + 1, ws.max_row + 1):
        name = ws.cell(r, name_col).value
        if isinstance(name, str) and name.strip():
            competitors.append({"name": name.strip(),
                                "results": [cell_code(ws.cell(r, first_col + i).value) for i in range(len(weeks) * 2)]})
    return weeks, competitors


def read_placar(ws, pins, week1):
    """Formato antigo (aba Placar)."""
    head_row, name_col, first_col = pins[0], pins[1] - 1, week1[1]
    weeks, col = [], first_col
    while True:
        label = ws.cell(week1[0], col).value
        if not label or (str(label).strip().upper() == "SEMANA" and col > first_col + 2):
            break
        weeks.append({"label": week_label(label, len(weeks) + 1),
                      "provas": [tc(ws.cell(head_row, col).value), tc(ws.cell(head_row, col + 1).value)]})
        col += 2
    competitors = []
    for r in range(head_row + 1, ws.max_row + 1):
        name = ws.cell(r, name_col).value
        if isinstance(name, str) and name.strip():
            competitors.append({"name": name.strip(),
                                "results": [cell_code(ws.cell(r, first_col + i).value) for i in range(len(weeks) * 2)]})
    return weeks, competitors


def read_scoring(wb):
    for ws in wb.worksheets:
        if "pontua" in ws.title.lower():
            out = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0] and isinstance(row[2] if len(row) > 2 else None, (int, float)):
                    out.append({"code": str(row[0]).strip().upper(), "label": str(row[1] or row[0]).strip(), "pts": row[2]})
            if out:
                return out
        tab = find(ws, r"TABELA DE PONTUA")
        if tab:
            out = []
            for r in range(tab[0] + 1, ws.max_row + 1):
                code, pts = ws.cell(r, tab[1] + 3).value, ws.cell(r, tab[1] + 4).value
                if code is not None and isinstance(pts, (int, float)):
                    label = str(ws.cell(r, tab[1]).value or code).strip().replace("EIMINADO", "ELIMINADO")
                    out.append({"code": str(code).strip().upper(), "label": label, "pts": pts})
            if out:
                return out
    return None


def read_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    weeks = competitors = None
    for ws in wb.worksheets:
        comp, w1, pins = find(ws, r"^COMPETIDOR(ES)?$"), find(ws, WEEK1), find(ws, r"^PINS$")
        if comp and w1 and w1[0] == comp[0] - 1 and not pins:
            weeks, competitors = read_model(ws, comp, w1)
            break
        if pins and w1:
            weeks, competitors = read_placar(ws, pins, w1)
            break
    if weeks is None:
        raise ValueError(f"{path.name}: formato não reconhecido (esperava o modelo MC ou a aba Placar)")
    code = code_for(re.sub(r"^\s*tabela\s+", "", path.stem, flags=re.I))
    return {"id": code, "title": title_for(code), "source": path.name, "weeks": weeks,
            "scoring": read_scoring(wb), "competitors": competitors}


def load_bundle():
    if not DATA_JS.exists():
        return {"default": None, "editions": []}
    text = DATA_JS.read_text(encoding="utf-8")
    obj = json.loads(text[text.index("=") + 1:].strip().rstrip(";"))
    if "editions" not in obj:  # formato bem antigo: uma edição só
        obj = {"default": "MC-2026", "editions": [{"id": "MC-2026", "title": "MasterChef 2026", **obj}]}
    remap = {}
    for e in obj["editions"]:  # ids antigos (masterchef-2026) viram MC-2026
        new = code_for(e["id"] if re.match(r"^MC-", e["id"], re.I) else e.get("title") or e["id"])
        remap[e["id"]] = new
        e["id"] = new
    obj["default"] = remap.get(obj.get("default"), obj.get("default"))
    return obj


def fill_from(ed, old):
    """Completa nomes de provas e pontuação que a planilha não trouxe, a partir da versão anterior."""
    if not old:
        return ed
    for i, w in enumerate(ed["weeks"]):
        if i < len(old["weeks"]):
            old_label = old["weeks"][i]["label"]
            if re.fullmatch(r"Semana \d+", w["label"]) and old_label and not re.fullmatch(r"Semana \d+", old_label):
                w["label"] = old_label  # ex.: mantém "Repescagem"
            for j in (0, 1):
                if not w["provas"][j] and old["weeks"][i]["provas"][j]:
                    w["provas"][j] = old["weeks"][i]["provas"][j]
    if not ed["scoring"]:
        ed["scoring"] = old.get("scoring")
    ed["title"] = old.get("title") or ed["title"]
    return ed


def main():
    files = [Path(a) for a in sys.argv[1:]] or sorted((ROOT / "Dados").glob("*.xls*"), key=lambda p: p.stat().st_mtime)
    files = [f for f in files if not f.name.startswith("~$")]
    bundle = load_bundle()
    by_id = {e["id"]: e for e in bundle["editions"]}
    fallback_scoring = next((e["scoring"] for e in bundle["editions"] if e.get("scoring")), None)
    for f in files:  # em ordem de modificação: a mais recente vence quando o código se repete
        ed = fill_from(read_workbook(f), by_id.get(code_for(re.sub(r"^\s*tabela\s+", "", f.stem, flags=re.I))))
        ed["scoring"] = ed["scoring"] or fallback_scoring
        if ed["id"] in by_id and by_id[ed["id"]].get("source") not in (None, f.name):
            print(f"  (aviso) {f.name} substitui {by_id[ed['id']].get('source')} na edição {ed['id']}")
        by_id[ed["id"]] = ed
        print(f"{f.name}: {ed['id']} — {len(ed['competitors'])} competidores, {len(ed['weeks'])} episódios")
    order = list(dict.fromkeys([e["id"] for e in bundle["editions"]] + list(by_id)))
    bundle["editions"] = [by_id[i] for i in order]
    if not bundle.get("default") or bundle["default"] not in by_id:
        bundle["default"] = bundle["editions"][-1]["id"] if bundle["editions"] else None
    DATA_JS.write_text(
        "// Gerado por scripts/extrair_dados.py (ou exportado pelo app). Cópia de segurança de todas as edições.\n"
        "window.MASTERCHEF_EDITIONS = " + json.dumps(bundle, ensure_ascii=False, indent=1) + ";\n",
        encoding="utf-8",
    )
    print(f"data.js atualizado: {len(bundle['editions'])} edição(ões), padrão = {bundle['default']}")


if __name__ == "__main__":
    main()
